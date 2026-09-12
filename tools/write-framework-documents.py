"""Build editable Word copies from the same reviewed data used by the game.

No internal original is opened or modified. Output is the parent project folder.
"""
from pathlib import Path
import json
import sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT.parent
F = json.loads((ROOT / 'content/competency-framework.json').read_text(encoding='utf-8'))
C = json.loads((ROOT / 'content/career-profiles.json').read_text(encoding='utf-8'))
NAVY='173650'; TEAL='32655E'; GREY='53636B'; PALE='EBF1EE'; GOLD='AD7A38'

def loc(x, lang):
    return x.get(lang, '') if isinstance(x, dict) else str(x)

def pair(lang, de, en): return de if lang=='de' else en

def shade(cell, color):
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),color);cell._tc.get_or_add_tcPr().append(sh)

def no_split(row):
    el=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(el)

def hyperlink(p, title, url):
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),TEAL);props.append(color)
    underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');props.append(underline);r.append(props)
    t=OxmlElement('w:t');t.text=title;r.append(t);link.append(r);p._p.append(link)

def table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,h in enumerate(headers):
        cell=t.rows[0].cells[i];cell.text=h;shade(cell,NAVY)
        for r in cell.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor.from_string('FFFFFF')
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for n,items in enumerate(rows):
        row=t.add_row();no_split(row)
        for i,value in enumerate(items):
            cell=row.cells[i];cell.text=str(value);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            if n%2==0:shade(cell,PALE)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
                for r in p.runs:r.font.size=Pt(9)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths):row.cells[i].width=Cm(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def para(doc, text, style=None): return doc.add_paragraph(text,style)
def heading(doc,text,level=1):return doc.add_heading(text,level)

def init(lang,title,subtitle):
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(2);s.bottom_margin=Cm(1.9);s.left_margin=Cm(2.1);s.right_margin=Cm(2.1)
    s.header_distance=Cm(.8);s.footer_distance=Cm(.8)
    normal=d.styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(10)
    normal.font.color.rgb=RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
    normal.element.get_or_add_rPr().append(OxmlElement('w:lang'))
    normal.element.rPr[-1].set(qn('w:val'),'de-CH' if lang=='de' else 'en-GB')
    for name,size in [('Title',32),('Subtitle',14),('Heading 1',21),('Heading 2',14),('Heading 3',11)]:
        st=d.styles[name];st.font.name='Aptos Display' if name in ['Title','Heading 1'] else 'Aptos';st.font.size=Pt(size)
        st.font.color.rgb=RGBColor.from_string(TEAL if name=='Heading 2' else NAVY)
        st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(13);st.paragraph_format.space_after=Pt(8)
    header=s.header.paragraphs[0];header.text='KOMPETENZHAUS  /  PSYCHOLOGIE'
    header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor.from_string(GREY)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(pair(lang,'Entwurf · September 2026  |  ','Draft · September 2026  |  ')).font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    d.core_properties.title=title;d.core_properties.subject=subtitle
    d.core_properties.author='Kompetenzhaus project';d.core_properties.keywords='Psychologie, Kompetenzentwicklung, AI, Future Skills, Entwurf'
    para(d,pair(lang,'MODELL 2.2 · ARBEITSENTWURF','MODEL 2.2 · WORKING DRAFT'))
    heading(d,title,0);para(d,subtitle,'Subtitle')
    para(d,loc(F['status'],lang))
    para(d,pair(lang,'Fachwissenschaft, Future Skills und KI-Kompetenzen gemeinsam entwickeln.','Develop disciplinary expertise, future skills and AI competence together.'))
    para(d,pair(lang,'16 Kompetenzen · 48 Kriterien · 3 Aufgabenniveaus · 12 berufliche Perspektiven','16 competencies · 48 criteria · 3 task levels · 12 professional perspectives'))
    para(d,pair(lang,'Stand: 12. September 2026. Für das Lernspiel und als Gesprächsgrundlage zur Curriculumentwicklung.','Version date: 12 September 2026. For the learning game and as a basis for curriculum discussions.'))
    para(d,pair(lang,'Die bestehenden Modulbeschreibungen bilden den Ausgangspunkt. Neue Lernziele und Aufgaben sind Vorschläge. Die ursprünglichen Modelle und Quelldokumente bleiben unverändert.','Existing module descriptions are the starting point. New objectives and tasks are proposals. The original models and source documents remain unchanged.'))
    return d

def page_heading(doc,text):
    p=heading(doc,text);p.paragraph_format.page_break_before=True;return p

def ai_phases(doc, lang, heading_level=1, new_page=False):
    across=F['aiAcrossCurriculum']
    title=heading(doc,pair(lang,'Ohne KI, mit KI und über KI lernen und prüfen','Learning and assessment without, with and about AI'),heading_level)
    if new_page:title.paragraph_format.page_break_before=True
    para(doc,loc(across['assessmentNotice'],lang))
    for phase in across['phasePolicy']['phases']:
        heading(doc,loc(phase['title'],lang),2)
        for field in ['purpose','learnerAction','assessmentUse']:para(doc,loc(phase[field],lang))
    heading(doc,pair(lang,'Bedingungen für Lern- und Prüfphasen','Conditions for learning and assessment phases'),2)
    for rule in across['phasePolicy']['assessmentRules']:para(doc,loc(rule['text'],lang),'List Bullet')


def ai_crosscurricular_pages(doc,lang):
    across=F['aiAcrossCurriculum']
    page_heading(doc,loc(across['title'],lang));para(doc,loc(across['principle'],lang))
    # Editable navigation cards; no ranking or numerical assessment plot.
    cards=doc.add_table(rows=3,cols=2);cards.autofit=False
    for index,row in enumerate(across['domainRows']):
        cell=cards.rows[index//2].cells[index%2];cell.width=Cm(8.4);shade(cell,PALE if index%2==0 else 'EDF2F1')
        domain=next(x for x in F['domains'] if x['id']==row['domainId'])
        p=cell.paragraphs[0];r=p.add_run(row['domainId']+'  '+loc(domain['title'],lang));r.bold=True
        cell.add_paragraph(' · '.join(row['competencyIds']))
        cell.add_paragraph(loc(row['summary'],lang))
        for p in cell.paragraphs:
            p.paragraph_format.space_after=Pt(6)
            for r in p.runs:r.font.size=Pt(9)
    for row in cards.rows:no_split(row)
    para(doc,loc(across['flowRule'],lang))
    page_heading(doc,pair(lang,'Wie die Bereiche in einer Fachaufgabe zusammenwirken','How domains connect within a disciplinary task'))
    table(doc,[pair(lang,'Verbindung','Connection'),pair(lang,'Übergang und kanonische Kriterien','Transition and canonical criteria')],[(x['fromDomainId']+' → '+x['toDomainId'],loc(x['label'],lang)+'\n'+', '.join(x['criterionIds'])) for x in across['flowEdges']],[2.5,14.3])
    heading(doc,pair(lang,'Wissen, Fähigkeiten, Haltungen und Werte','Knowledge, capabilities, attitudes and values'),2)
    for dimension in across['dimensions']:
        p=para(doc,'');p.add_run(loc(dimension['title'],lang)+': ').bold=True;p.add_run(loc(dimension['description'],lang))
    page_heading(doc,pair(lang,'Woran die menschliche Leistung erkennbar wird','What makes the human contribution visible'))
    first_domain_paragraph=len(doc.paragraphs)
    for row in across['domainRows']:
        heading(doc,row['domainId']+' · '+', '.join(row['competencyIds']),2)
        para(doc,loc(row['learnerResponsibility'],lang));para(doc,loc(row['evidenceFocus'],lang))
        p=para(doc,pair(lang,'Kriterien: ','Criteria: ')+', '.join(row['criterionIds']));p.paragraph_format.space_after=Pt(10)
    for p in doc.paragraphs[first_domain_paragraph:]:
        p.paragraph_format.space_before=Pt(5 if p.style.name.startswith('Heading') else 0)
        p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.line_spacing=1.02
        if not p.style.name.startswith('Heading'):
            for r in p.runs:r.font.size=Pt(9.5)
    ai_phases(doc,lang,new_page=True)
    for example in across['examples']:
        page_heading(doc,loc(example['title'],lang))
        para(doc,pair(lang,'Mögliche Anschlüsse: ','Possible module links: ')+', '.join(example['moduleIds']))
        para(doc,pair(lang,'Kriterien: ','Criteria: ')+', '.join(example['criterionIds']))
        para(doc,pair(lang,'Aufgabenbeispiel im Entwurf. Die Modulnennung ist eine Lerngelegenheit und keine verbindliche Vorgabe.','Proposed task example. The named modules indicate opportunities, not binding requirements.'))
        for phase in example['phases']:
            title=next(x['title'] for x in across['phasePolicy']['phases'] if x['contextId']==phase['contextId'])
            heading(doc,loc(title,lang),2);para(doc,loc(phase['activity'],lang))
            p=para(doc,'');p.add_run(pair(lang,'Vorgeschlagener Nachweis: ','Proposed evidence: ')).bold=True;p.add_run(loc(phase['evidence'],lang))
        heading(doc,pair(lang,'Beurteilungsvorschlag','Assessment proposal'),2);para(doc,loc(example['assessmentProposal'],lang))


def build(lang):
    title=pair(lang,'Psychologische Handlungskompetenz','Competence in psychological practice and research')
    d=init(lang,title,pair(lang,'Ein zukunftsgerichtetes Kompetenzaufbaumodell für das Kompetenzhaus','A future-oriented competence development model for the Kompetenzhaus'))
    heading(d,pair(lang,'Die Entscheidungen hinter dem Modell','The decisions behind the model'),2)
    points=[
      pair(lang,'Kompetenzen beschreiben begründetes Handeln. Wissen über psychologische Theorien, Methoden und Befunde bleibt dafür unverzichtbar.','Competencies describe reasoned action. Knowledge of psychological theories, methods and findings remains essential.'),
      loc(F['designPrinciple'],lang),
      pair(lang,'Future Skills und KI werden in Aufgaben sichtbar: eigene Probleme formulieren, Alternativen entwickeln, mit Unsicherheit arbeiten, Systeme gestalten und Entscheidungen verantworten.','Future skills and AI become visible in tasks: formulate problems, develop alternatives, work with uncertainty, design systems and take responsibility for decisions.'),
      pair(lang,'Drei Niveaus beschreiben die Anforderungen einer Aufgabe. Vier Studienmeilensteine ordnen mögliche Lerngelegenheiten. Beides wird getrennt geführt.','Three levels describe task demands. Four study milestones organise possible learning opportunities. They are recorded separately.'),
      pair(lang,'Ein Berufsprofil zeigt Tätigkeiten und nächste Lernmöglichkeiten. Es liefert keine Prozentzahl zur persönlichen Eignung.','A professional profile shows activities and next learning opportunities. It does not produce a percentage of personal suitability.')]
    for text in points:para(d,text,'List Bullet')
    page_heading(d,pair(lang,'Das Modell auf einen Blick','The model at a glance'))
    table(d,[pair(lang,'Bereich','Domain'),pair(lang,'Kompetenzen','Competencies')],[(loc(dm['title'],lang),'\n'.join(c['id']+'  '+loc(c['title'],lang) for c in F['competencies'] if c['domainId']==dm['id'])) for dm in F['domains']],[5.1,11.7])
    para(d,loc(F['limitation'],lang))
    heading(d,pair(lang,'Was eine Zuordnung leistet','What the classification does'),2)
    para(d,pair(lang,'Die Trennung erfolgt nach der beurteilten Leistung. Bei einem KI-gestützten Gutachten können mehrere Kriterien relevant sein: fachliche Belege prüfen (R4), einen funktionierenden Ablauf testen (T3), Schutzinteressen abwägen (V1) und das Ergebnis verständlich darstellen (K1). Jedes Kriterium wird separat beurteilt; dieselbe Leistung erhält keinen zweiten KI- oder Future-Skills-Bonus.','The classification follows the performance being judged. An AI-assisted report may involve several criteria: checking evidence (R4), testing a functional workflow (T3), weighing protected interests (V1) and communicating clearly (K1). Each criterion is judged separately; the same performance receives no second AI or future-skills bonus.'))
    heading(d,pair(lang,'Wissenschaft und Zukunft zusammenhalten','Keep science and future capabilities together'),2)
    para(d,pair(lang,'Ein neues Werkzeug darf die fachliche Begründung nicht ersetzen. Die Aufgaben verbinden deshalb Wissen mit Entscheidungen, Gestaltung und Rückmeldung. Ein Masterauftrag kann anspruchsvoller sein, weil die Frage offen ist, Befunde widersprechen oder Verantwortung geteilt werden muss. Ein höherer KI-Anteil allein macht ihn nicht anspruchsvoller.','A new tool must not replace disciplinary reasoning. Tasks therefore connect knowledge with decisions, design and feedback. A Master’s task may be more demanding because its question is open, findings conflict or responsibility is shared. A larger AI component alone does not increase its level.'))
    ai_crosscurricular_pages(d,lang)
    page_heading(d,pair(lang,'Spiralcurriculum und Kompetenzstufen','Spiral curriculum and competence levels'))
    table(d,[pair(lang,'Niveau','Level'),pair(lang,'Anforderung','Demand')],[(x['label']+' · '+loc(x['title'],lang),loc(x['description'],lang)) for x in F['levelDefinitions']],[4.9,11.9])
    para(d,loc(F['milestoneRule'],lang))
    para(d,pair(lang,'Ein späterer Auftrag kann dasselbe Kriterium unter schwierigeren Bedingungen aufgreifen. Die Stufen werden nicht addiert. Ein einzelner gelungener Auftrag belegt keine allgemeine Beherrschung in allen Situationen. Je nach Vorwissen können anspruchsvolle Aufgaben früher und Grundlagenübungen später sinnvoll sein.','A later task can revisit the same criterion under more demanding conditions. Levels are not added together. Success on one task does not establish general mastery across situations. Prior experience may justify advanced tasks earlier or foundational practice later.'))
    heading(d,pair(lang,'Beispiel: KI-Ethik I–III','Example: AI ethics I–III'),2)
    table(d,[pair(lang,'Zielniveau und Gelegenheit','Target level and opportunity'),pair(lang,'Aufgabe','Task')],[(str(x['level'])+' · '+', '.join(x['moduleIds']),loc(x['task'],lang)) for x in F['spiralExample']['occasions']],[4.9,11.9])
    para(d,loc(F['spiralExample']['note'],lang))
    heading(d,pair(lang,'Vier Studienmeilensteine','Four study milestones'),2)
    for x in F['studyMilestones']:para(d,loc(x['title'],lang)+' · '+pair(lang,'typische Semester ','typical semesters ')+x['typicalSemesters'],'List Bullet')
    for c in F['competencies']:
        d.add_page_break();para(d,loc(next(dm for dm in F['domains'] if dm['id']==c['domainId'])['title'],lang).upper())
        heading(d,c['id']+'  '+loc(c['title'],lang));para(d,loc(c['scope'],lang))
        heading(d,pair(lang,'Beobachtbare Kriterien','Observable criteria'),2)
        for k in c['criteria']:para(d,k['id']+'  '+loc(k['text'],lang),'List Bullet')
        heading(d,pair(lang,'Progression in Aufgaben','Progression through tasks'),2)
        table(d,[pair(lang,'Niveau','Level'),pair(lang,'Konkretisierung','Description')],[(str(x['level']),loc(x['description'],lang)) for x in c['levels']],[1.7,15.1])
        heading(d,pair(lang,'Abgrenzung','Boundary'),2);para(d,loc(c['boundary'],lang))
        para(d,pair(lang,'Die genannten Niveaus sind Zielanforderungen. Eine Beurteilung braucht eine konkrete Aufgabe, Bedingungen, Kriterien und Rückmeldung.','These levels are target demands. Assessment requires a specific task, conditions, criteria and feedback.'))
        d.add_page_break()
        para(d,c['id']+'  /  '+pair(lang,'KI UND FUTURE SKILLS IN DER FACHAUFGABE','AI AND FUTURE SKILLS IN DISCIPLINARY TASKS'))
        heading(d,loc(c['title'],lang))
        for field,de,en in [('possibleRole','Welche Rolle KI spielen kann','The possible role of AI'),('learnerResponsibility','Was Studierende selbst leisten','The student’s own contribution'),('assessmentFocus','Was beurteilt wird','What is assessed'),('exampleTask','Beispiel einer aktivierenden Lernaufgabe','Example of an active learning task')]:
            heading(d,pair(lang,de,en),2);para(d,loc(c['aiIntegration'][field],lang))
        para(d,pair(lang,'Kriterienbezüge: ','Criterion references: ')+', '.join(c['aiIntegration']['relatedCriterionIds']))
        lenses=[x for x in F['futureLenses'] if c['id'] in x['competencyIds']]
        if lenses:
            heading(d,pair(lang,'Future-Skills-Perspektiven','Future skills perspectives'),2)
            para(d,' · '.join(loc(x['title'],lang) for x in lenses))
        para(d,pair(lang,'Die Aufgabe ist ein Vorschlag und kann mit vorbereiteten KI-Ausgaben oder einer Simulation bearbeitet werden. Sie setzt keinen kostenpflichtigen KI-Zugang voraus. Die drei Kontexte ohne KI, mit KI und über KI bleiben je nach Lernziel wählbar.','This proposed task can use prepared AI outputs or a simulation. It requires no paid AI access. The contexts without AI, with AI and about AI remain available according to the learning objective.'))
    page_heading(d,pair(lang,'Future Skills und KI ausdrücklich mitdenken','Make future skills and AI explicit'))
    para(d,loc(F['futureCoverageRule'],lang))
    for item in F['futureLenses']:
        heading(d,loc(item['title'],lang),2);para(d,loc(item['description'],lang));para(d,pair(lang,'Kriterien: ','Criteria: ')+', '.join(item['criterionIds']))
    heading(d,pair(lang,'Sechs Fragen an jede Zukunftsaufgabe','Six questions for every future-oriented task'),2)
    for item in F['futureDesignTests']:para(d,loc(item,lang),'List Bullet')
    page_heading(d,pair(lang,'Kompetenzkompasse für Tätigkeiten','Competence compasses for professional activities'))
    para(d,loc(C['notice'],lang))
    para(d,pair(lang,'Die gemeinsame Grundlage bleibt für alle sichtbar. Innerhalb eines Berufsprofils werden ausgewählte Kriterien nach beruflichen Tätigkeiten angeordnet. Ein Kriterium erscheint dort höchstens einmal. Unterschiedliche Berufsansichten sind verschiedene Perspektiven auf denselben Lernweg; ihre Werte werden nicht zusammengezählt.','The shared foundation remains visible. Within a professional profile, selected criteria are organised by professional activities. Each criterion appears at most once within that profile. Different profiles are views of the same learning pathway; their values are not added together.'))
    for role in C['roles']:
        heading(d,loc(role['title'],lang),2)
        for index,facet in enumerate(role['facets']):
            p=para(d,'');p.add_run(loc(facet['title'],lang)+': ').bold=True
            context=pair(lang,'Optionaler digitaler/KI-Kontext. ','Optional digital/AI context. ') if facet.get('optionalContext') else ''
            p.add_run(context+loc(facet['activity'],lang)+' ['+', '.join(facet['criterionIds'])+']')
            p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.02
            p.paragraph_format.keep_with_next=index==len(role['facets'])-1
            for run in p.runs:run.font.size=Pt(9.5)
        p=para(d,'');p.paragraph_format.space_after=Pt(4)
        hyperlink(p,pair(lang,'Fachliche Orientierung der Tätigkeitsauswahl','Source informing the activity selection'),role['sourceUrl'])
    page_heading(d,pair(lang,'Lerngelegenheiten, Übung und Nachweise','Opportunities, practice and evidence'))
    table(d,[pair(lang,'Ebene','Layer'),pair(lang,'Was angezeigt werden darf','What can be shown')],[
      (pair(lang,'Geplant','Planned'),pair(lang,'Ein Modul bietet ein Lernziel auf einem Zielniveau. Das ist eine Gelegenheit.','A module offers an objective at a target level. This is an opportunity.')),
      (pair(lang,'Im Spiel geübt','Practised in the game'),pair(lang,'Ein Lerncheck oder eine Quest wurde bearbeitet. Das sagt nicht, dass alle Modulziele beherrscht werden.','A check or quest was completed. This does not show mastery of all module objectives.')),
      (pair(lang,'Selbst eingeschätzt','Self-reported'),pair(lang,'Eine Person beschreibt ihre Erfahrung oder Zuversicht. Das bleibt ein gekennzeichnetes Selbsturteil.','A person describes experience or confidence. This remains a labelled self-report.')),
      (pair(lang,'Kriterienbezogen beurteilt','Assessed against criteria'),pair(lang,'Eine konkrete Leistung wurde unter bekannten Bedingungen anhand benannter Kriterien beurteilt. Das aktuelle Spiel besitzt keine solchen Hochschulnachweise.','A specific performance was assessed under known conditions against named criteria. The current game contains no such university assessment records.'))
    ],[4.2,12.6])
    para(d,pair(lang,'Für einen Nachweis braucht es Kriterium, Zielniveau, Aufgabe, Kontext, Belegart, Urteil und Rückmeldung. Wiederholungen ergänzen die Evidenzgeschichte. Sie erzeugen keine zweite Kompetenz. Nicht beurteilt ist ein eigener Zustand und nicht der Wert null.','An evidence record needs a criterion, target level, task, context, evidence type, judgement and feedback. Repetitions extend the evidence history. They do not create another competence. Unassessed is a separate state, not a score of zero.'))
    heading(d,pair(lang,'Übertragung in das Spiel','Application in the game'),2)
    for text in [
      pair(lang,'Die Modulübersicht zeigt den übernommenen Ausgangsstand und den künftigen Lernentwurf getrennt. Auswahlmodule ersetzen die Ziele ihres Slots; sie werden nicht zusätzlich gezählt.','The module overview separates the inherited baseline from the future learning design. Electives replace their slot’s objectives rather than being added again.'),
      pair(lang,'Das Profil zeigt die Lerngelegenheiten der gebauten Module mit Zielniveaus und passenden Aufgaben. Hausbau und Quizquoten werden nicht in persönliche Kompetenzwerte umgerechnet.','The profile shows learning opportunities in built modules, with target levels and relevant tasks. Construction and quiz scores are not converted into personal competence levels.'),
      pair(lang,'Der Berufskompass bietet konkrete Tätigkeiten und nächste mögliche Module. Eine allgemeine Passungszahl und eine vermeintliche berufliche Zielkurve entfallen.','The professional compass offers concrete activities and possible next modules. It omits a general fit score and a purported professional target curve.'),
      pair(lang,'Gestaltung, Erkundung und Quests schaffen Anlass für Entscheidungen und Rückmeldung. Kosmetische Belohnungen stehen getrennt von fachlichen Nachweisen.','Design, exploration and quests create opportunities for decisions and feedback. Cosmetic rewards remain separate from disciplinary evidence.')]:para(d,text,'List Bullet')
    heading(d,pair(lang,'Fachlich prüfen und erproben','Review and pilot'),2)
    para(d,pair(lang,'Vor curricularer Verwendung sollen Lehrende die Fachabdeckung und Aufgabenanforderungen prüfen, Studierende die Verständlichkeit und Arbeitsbelastung erproben und Berufspersonen die Tätigkeitsansichten beurteilen. Beispielarbeiten dienen anschliessend dazu, Kriteriengrenzen und Beurteilungsübereinstimmung zu testen. Dieses Dokument behauptet keine bereits erfolgte Zustimmung oder empirische Validierung.','Before curricular use, lecturers should review disciplinary coverage and task demands, students should test clarity and workload, and practitioners should examine the activity views. Sample work can then be used to test criterion boundaries and agreement between assessors. This document does not claim completed approval or empirical validation.'))
    page_heading(d,pair(lang,'Zuordnung bisheriger Begriffe','Mapping earlier terms'))
    para(d,pair(lang,'Die Zuordnung erklärt die begriffliche Weiterentwicklung. Alte Punktwerte werden nicht übertragen: Die bisherige Berechnung anhand gebauter ECTS eignet sich nicht als Beleg für die neuen Kriterien.','This mapping explains the conceptual revision. Old scores are not transferred: the previous calculation based on built credits cannot establish evidence for the new criteria.'))
    table(d,[pair(lang,'Bisherige ID','Previous ID'),pair(lang,'Kanonische Zuordnung','Canonical mapping')],[(x['legacyId'],', '.join(x['competencyIds'])) for x in F['legacyCrosswalk']],[4.2,12.6])
    heading(d,pair(lang,'Rahmen werden zu Quellen, nicht zu parallelen Skalen','Frameworks become sources, not parallel scales'),2)
    para(d,pair(lang,'Der St. Galler Rahmen unterscheidet operative und reflexive KI-Kompetenzen und verarbeitet bereits weitere Modelle, darunter AIComp. AIComp ordnet Zukunftskompetenzen in zwölf Felder. UNESCO liefert Perspektiven auf Menschenorientierung, Ethik und technische Gestaltung. Der UZH-/ETH-Rahmen ergänzt die besondere Rolle von Dozierenden. Psychologische Fachwissenschaft und berufliche Tätigkeiten werden im vorliegenden Entwurf ausdrücklich ergänzt. Die Zuordnungen sind eine eigene didaktische Synthese.','The St. Gallen framework distinguishes operational and reflective AI competencies and already draws on other models, including AIComp. AIComp organises future capabilities into twelve fields. UNESCO contributes perspectives on human agency, ethics and technical design. The UZH/ETH framework adds the specific lecturer role. This draft explicitly retains psychological science and professional activities. These mappings are an original educational synthesis.'))
    page_heading(d,pair(lang,'Quellen und Geltungsgrenzen','Sources and scope'))
    for s in F['sources']:
        p=para(d,'');hyperlink(p,s['title'],s['url'])
        if lang=='en':para(d,s['role'])
    heading(d,pair(lang,'Lokale Ausgangspunkte','Local starting points'),2)
    para(d,pair(lang,'Die zusätzlich bereitgestellte AIComp-Foliensammlung trägt im Dokument den Stand 17.11.2023. Für diese Revision wurden insbesondere Systemgestaltung (S. 11–13), aktive Steuerung und Selbstbestimmung (S. 29–34) sowie Kommunikation über KI mit Menschen (S. 41–43) herangezogen. Die psychologischen Aufgaben und Kriterienzuordnungen sind eigene Ausarbeitungen.','The additionally supplied AIComp slide deck is dated 17 November 2023 within the document. This revision particularly draws on system design (pp. 11–13), active control and self-determination (pp. 29–34), and communication about AI with people (pp. 41–43). The psychology tasks and criterion mappings are original applications.'))
    para(d,pair(lang,'Zusätzlich berücksichtigt wurden das bestehende Kompetenzaufbaumodell des Psychologischen Instituts, der KI-im-Curriculum-Kompass UZH, die bereitgestellten September-Unterlagen und die Benchmark-Ausgaben. Diese Unterlagen wurden ausschliesslich gelesen. Ihre Rohdateien werden nicht in die öffentliche Spiel-Repository übernommen. Der Kompass ist eine Planungshilfe; seine acht Entwicklungsperspektiven werden nicht als acht Studierendenkompetenzen behandelt.','The existing Psychology Institute competence model, the UZH AI-in-curriculum compass, the supplied September materials and benchmark outputs also informed the work. These sources were read only. Their raw files are not included in the public game repository. The compass supports planning; its eight development perspectives are not treated as eight student competencies.'))
    para(d,pair(lang,'Die publizierten Spielinhalte sind die technische Ausgangsbasis für Modultitel, Codes und Credits. Ihre Übernahme ist keine erneute Verifikation der geltenden Studienordnung. Zukünftige Modulbeschreibungen bleiben separat gekennzeichnete Vorschläge.','The published game content is the technical baseline for module titles, codes and credits. Reusing it does not constitute renewed verification of current study regulations. Future module descriptions remain separately labelled proposals.'))
    para(d,pair(lang,'Die drei Aufgabenniveaus sind eine begründete Designentscheidung dieses Entwurfs. Sie sind weder aus Studiensemestern abgeleitet noch als psychometrische Skala validiert.','The three task levels are a reasoned design choice in this draft. They are neither derived from semesters nor validated as a psychometric scale.'))
    import hashlib
    para(d,'Framework: '+F['version'])
    para(d,'SHA256 content/competency-framework.json: '+hashlib.sha256((ROOT/'content/competency-framework.json').read_bytes()).hexdigest())
    out=OUT/('Kompetenzaufbaumodell-Psychologie-Zukunft-v2.2-'+lang.upper()+'.docx')
    d.save(out)
    return out

if __name__=='__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    for language in ['de','en']:
        path=build(language);print(path)
