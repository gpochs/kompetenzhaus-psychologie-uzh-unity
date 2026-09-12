"""Create two editable reference volumes from the final module-design source.

Uses the shared framework-document layout helpers without changing their file.
Original descriptions and authored proposals remain visibly separate. No Word,
PDF renderer, network connection, paid service or original-source edit is used.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "content/module-learning-design.json"
FRAMEWORK = ROOT / "content/competency-framework.json"
EXPECTED_VERSION = "2.1.0-draft"
EXPECTED_ENTRIES = 46
EXPECTED_OBJECTIVES = 153
ROMAN = {1: "I", 2: "II", 3: "III"}


def load_layout():
    spec = importlib.util.spec_from_file_location("framework_document_layout", ROOT / "tools/write-framework-documents.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def fail(message):
    raise ValueError(message)


def local(value, language):
    if not isinstance(value, dict) or not isinstance(value.get(language), str) or not value[language].strip():
        fail(f"Missing non-empty {language} source text: {value!r}")
    return value[language]


def validate_source(data, framework):
    if data.get("frameworkVersion") != EXPECTED_VERSION or framework.get("version") != EXPECTED_VERSION:
        fail("Both source files must be final framework version 2.1.0-draft.")
    if data.get("schema") != "kompetenzhaus.module-learning-design" or data.get("schemaVersion") != 1 or data.get("status") != "design-proposal":
        fail("Unsupported module learning design or missing proposal status.")
    modules = data.get("modules", [])
    if len(modules) != EXPECTED_ENTRIES or sum(len(m["proposal"]["objectives"]) for m in modules) != EXPECTED_OBJECTIVES:
        fail("Expected 46 entries and 153 authored objectives in the final source.")
    if len({m["id"] for m in modules}) != len(modules):
        fail("Duplicate module id.")
    ids = {m["id"] for m in modules}
    criteria = {c["id"] for competency in framework["competencies"] for c in competency["criteria"]}
    contexts = {c["id"] for c in framework["contexts"]}
    objectives = set()
    for source in data.get("sources", []):
        if not source.get("path") or not source.get("sha256"):
            continue
        path = (ROOT / source["path"]).resolve()
        if not path.is_relative_to(ROOT.resolve()) or not path.is_file() or sha256(path).lower() != source["sha256"].lower():
            fail(f"Source provenance mismatch: {source['path']}")
    for module in modules:
        baseline, proposal = module["baseline"], module["proposal"]
        if module["kind"] not in ("slot", "option") or baseline.get("status") != "inherited-public-game-current-baseline" or proposal.get("status") != "design-proposal":
            fail(f"Unexpected module status: {module['id']}")
        if baseline.get("id") != module["id"] or baseline.get("code") != module["code"]:
            fail(f"Baseline identity differs from module identity: {module['id']}")
        for language in ("de", "en"):
            for key in ("title", "description"):
                local(baseline[key], language)
            for key in ("knowledgeAnchors", "futureSummary", "activitySequence", "aiRole", "independentEvidence", "assessmentProposal", "workloadIntegration", "openDecisions"):
                local(proposal[key], language)
        for objective in proposal["objectives"]:
            if objective["id"] in objectives or objective["criterionId"] not in criteria or objective["targetLevel"] not in ROMAN:
                fail(f"Invalid objective identity, criterion or level: {objective['id']}")
            objectives.add(objective["id"])
            if not objective["contextIds"] or not set(objective["contextIds"]).issubset(contexts):
                fail(f"Unknown or absent context: {objective['id']}")
            for language in ("de", "en"):
                for key in ("text", "task", "evidence", "successCriteria"):
                    local(objective[key], language)
        for direction in ("prior", "next"):
            for link in proposal["spiralLinks"][direction]:
                if link["moduleId"] not in ids:
                    fail(f"Unresolved spiral link: {link['moduleId']}")
                for language in ("de", "en"):
                    local(link["reason"], language)
    if len(objectives) != EXPECTED_OBJECTIVES:
        fail("Objective completeness check failed.")


def bookmark(paragraph, name, index):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(index))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(index))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def internal_link(paragraph, text, anchor, color):
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    link.set(qn("w:history"), "1")
    run, props = OxmlElement("w:r"), OxmlElement("w:rPr")
    shade = OxmlElement("w:color")
    shade.set(qn("w:val"), color)
    props.append(shade)
    run.append(props)
    content = OxmlElement("w:t")
    content.text = text
    content.set(qn("xml:space"), "preserve")
    run.append(content)
    link.append(run)
    paragraph._p.append(link)


def add_page_reference(paragraph, anchor):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f"PAGEREF {anchor} \\h")
    field.set(qn("w:dirty"), "true")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "…"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def normalised(text):
    return re.sub(r"\s+", " ", text).strip()


def document_text(document):
    # Word hyperlinks and fields are included, unlike older paragraph.text APIs.
    return " ".join(el.text or "" for el in document.element.body.iter(qn("w:t")))


def verify_output(path, required, module_ids, objective_ids):
    doc = Document(path)
    text = normalised(document_text(doc))
    missing = [label for label, value in required if normalised(value) not in text]
    if missing:
        fail("Missing source fields after DOCX round trip: " + ", ".join(missing[:15]))
    for id in module_ids:
        if text.count(f"[module:{id}]") != 1:
            fail(f"Module marker missing or duplicated: {id}")
    for id in objective_ids:
        if text.count(f"[objective:{id}]") != 1:
            fail(f"Objective marker missing or duplicated: {id}")
    settings = doc.settings.element
    if settings.find(qn("w:documentProtection")) is not None:
        fail("Output unexpectedly contains editing protection.")
    if list(doc.element.body.iter(qn("w:drawing"))):
        fail("Expected editable text; unexpected drawing objects found.")
    return {"path": str(path), "words": len(text.split()), "modules": len(module_ids), "objectives": len(objective_ids), "verifiedSourceFields": len(required), "sha256": sha256(path)}


def build(language, data, framework, layout, output_folder):
    choose = lambda de, en: de if language == "de" else en
    text = lambda value: local(value, language)
    title = choose("Modulbeschreibungen\nPsychologie der Zukunft", "Module descriptions\nPsychology for the future")
    subtitle = choose("Lernentwürfe für Bachelor und Master · Band DE · Version 2.1", "Learning designs for Bachelor and Master study · English volume · Version 2.1")
    document = layout.init(language, title, subtitle)
    # Retain the shared page setup, styles, header and footer. Author this volume's
    # cover from module data rather than inheriting the framework-volume cover.
    for child in list(document.element.body):
        if child.tag != qn("w:sectPr"):
            document.element.body.remove(child)
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.04
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True
    for name, size in (("Heading 1", 20), ("Heading 2", 14), ("Heading 3", 11.5)):
        document.styles[name].font.size = Pt(size)
        document.styles[name].paragraph_format.space_before = Pt(9)
        document.styles[name].paragraph_format.space_after = Pt(6)
    document.core_properties.title = title.replace("\n", ": ")
    document.core_properties.subject = subtitle
    document.core_properties.comments = f"Editable design proposal. Source SHA256: {sha256(SOURCE)}"
    document.core_properties.language = "de-CH" if language == "de" else "en-GB"
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    document.settings.element.append(update)
    required = []

    def paragraph(value, style=None):
        return layout.para(document, value, style)

    def source_text(label, value, style=None):
        translated = text(value)
        required.append((label, translated))
        return paragraph(translated, style)

    def heading(value, level=1):
        return layout.heading(document, value, level)

    def labelled(label, value, identifier=None):
        translated = text(value) if isinstance(value, dict) else str(value)
        if identifier:
            required.append((identifier, translated))
        p = paragraph("")
        p.add_run(label + " ").bold = True
        p.add_run(translated)
        return p

    paragraph(choose("MODULBAND 2.1 · ARBEITSENTWURF", "MODULE VOLUME 2.1 · WORKING DRAFT"))
    heading(title, 0)
    paragraph(subtitle, "Subtitle")
    paragraph(choose("43 Studienbausteine · 3 Wahlvarianten · 153 Lernziele", "43 curriculum slots · 3 elective alternatives · 153 learning objectives"))
    source_text("proposalNotice", data["proposalNotice"])
    paragraph(choose("Stand: 12. September 2026. Nachschlageband für das Kompetenzhaus und die fachliche Curriculumentwicklung.", "Version date: 12 September 2026. A reference volume for the Kompetenzhaus and disciplinary curriculum development."))
    paragraph(choose("Alle Beschreibungen sind editierbar. Die ursprünglichen Angaben sind als Ausgangsstand gekennzeichnet; neue Lernziele, Aufgaben und Prüfideen stehen im Entwurfsabschnitt.", "All descriptions are editable. Original module information is labelled as the baseline; new objectives, activities and assessment ideas appear in the proposal section."))

    heading(choose("So ist der Band aufgebaut", "How to use this volume")).paragraph_format.page_break_before = True
    source_text("sourceNotice", data["sourceNotice"])
    source_text("levelRule", data["levelRule"])
    source_text("costRule", data["costRule"])
    for key, de, en in (("evidenceRule", "Nachweise und Leistung", "Evidence and performance"), ("spiralRule", "Spiralförmige Entwicklung", "Spiral development"), ("aiEthicsRule", "KI und ethische Verantwortung", "AI and ethical responsibility"), ("toolSeparationRule", "Technisches Verstehen, Anwenden und Gestalten", "Technical understanding, application and design")):
        if data.get(key):
            heading(choose(de, en), 2)
            source_text(key, data[key])
    paragraph(choose("Bei einer Wahlvariante ersetzen deren Lernziele den Entwurf des zugehörigen Studienplatzes. Die 153 Ziele dieses Nachschlagebands umfassen alle Alternativen; sie werden nicht gemeinsam einem individuellen Studienweg zugerechnet.", "When an elective alternative is selected, its objectives replace the proposal for its curriculum slot. The 153 objectives in this reference volume include every alternative; they are not all attributed to one individual study pathway."))
    heading(choose("Drei Aufgabenniveaus", "Three task levels"), 2)
    for level in framework["levelDefinitions"]:
        labelled(f"{level['label']} · {text(level['title'])}:", level["description"], "level:" + str(level["level"]))

    modules = data["modules"]
    module_by_id = {module["id"]: module for module in modules}
    anchors = {module["id"]: f"Module_{index:03d}" for index, module in enumerate(modules, 1)}
    criteria = {criterion["id"]: criterion for competency in framework["competencies"] for criterion in competency["criteria"]}
    contexts = {context["id"]: text(context["title"]) for context in framework["contexts"]}
    milestones = {milestone["id"]: text(milestone["title"]) for milestone in framework["studyMilestones"]}
    heading(choose("Modulregister", "Module register")).paragraph_format.page_break_before = True
    paragraph(choose("Die Modultitel sind mit ihren Beschreibungen verknüpft. Die Seitenspalte wird beim Aktualisieren der Word-Felder gefüllt.", "Module titles link to their descriptions. The page column is populated when Word fields are updated."))
    headers = [choose("Code", "Code"), choose("Modul / Studienbaustein", "Module / curriculum slot"), "ECTS", choose("Typ", "Type"), choose("Seite", "Page")]
    rows = [(module["code"], text(module["baseline"]["title"]), str(module["baseline"]["ects"]), choose("Wahlvariante", "Alternative") if module["kind"] == "option" else choose("Studienplatz", "Slot"), "") for module in modules]
    register = layout.table(document, headers, rows, [3.0, 8.2, 1.1, 3.0, 1.5])
    for index, module in enumerate(modules, 1):
        row = register.rows[index]
        p = row.cells[1].paragraphs[0]
        p.clear()
        internal_link(p, text(module["baseline"]["title"]), anchors[module["id"]], layout.TEAL)
        page = row.cells[4].paragraphs[0]
        page.clear()
        add_page_reference(page, anchors[module["id"]])
    for row in register.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                for run in p.runs:
                    run.font.size = Pt(10.5)

    for index, module in enumerate(modules, 1):
        baseline, proposal = module["baseline"], module["proposal"]
        section = choose("WAHLVARIANTE", "ELECTIVE ALTERNATIVE") if module["kind"] == "option" else choose("STUDIENBAUSTEIN", "CURRICULUM SLOT")
        paragraph(f"{index:02d} / {len(modules)}  ·  {section}").paragraph_format.page_break_before = True
        title_paragraph = heading(text(baseline["title"]))
        bookmark(title_paragraph, anchors[module["id"]], index)
        required.append((module["id"] + ":title", text(baseline["title"])))
        code_line = paragraph(f"{module['code']}  ·  {baseline['ects']} ECTS")
        code_line.paragraph_format.keep_with_next = True
        marker = paragraph(f"[module:{module['id']}]")
        marker.runs[0].font.size = Pt(9)
        marker.runs[0].font.color.rgb = RGBColor.from_string(layout.GREY)
        heading(choose("1  Übernommener Ausgangsstand", "1  Inherited baseline"), 2)
        house = choose("Bachelor", "Bachelor") if baseline["houseId"] == "bsc" else "Master"
        semester_label = choose("Semester im Ausgangsspiel", "Source-game semester") if baseline["semester"] else choose("Semesterfeld im Ausgangsdatensatz", "Semester field in source data")
        paragraph(f"{house} · {semester_label} {baseline['semester']} · {milestones[baseline['stageId']]}")
        source_text(module["id"] + ":baseline", baseline["description"])
        if module.get("parentSlotIds"):
            labelled(choose("Wählbar im Studienplatz:", "Selectable in curriculum slot:"), "; ".join(f"{module_by_id[id]['code']} · {text(module_by_id[id]['baseline']['title'])}" for id in module["parentSlotIds"]))
        alternatives = [other for other in modules if module["id"] in other.get("parentSlotIds", [])]
        if alternatives:
            labelled(choose("Im Band enthaltene Wahlvarianten:", "Alternatives included in this volume:"), "; ".join(f"{other['code']} · {text(other['baseline']['title'])}" for other in alternatives))
        paragraph(choose("Übernommene Beschreibung aus dem öffentlichen Ausgangsspiel; keine erneute Bestätigung geltender Studienregeln.", "Description inherited from the public source game; not a renewed confirmation of current study regulations."))
        heading(choose("2  Entwurf der Lernumgebung", "2  Proposed learning environment"), 2)
        source_text(module["id"] + ":futureSummary", proposal["futureSummary"])
        labelled(choose("Fachliche Wissensgrundlage:", "Disciplinary knowledge base:"), proposal["knowledgeAnchors"], module["id"] + ":knowledgeAnchors")
        labelled(choose("Vorgeschlagener Ablauf:", "Proposed activity sequence:"), proposal["activitySequence"], module["id"] + ":activitySequence")
        labelled(choose("KI-Rolle:", "Role of AI:"), proposal["aiRole"], module["id"] + ":aiRole")
        labelled(choose("Eigenleistungsnachweis:", "Evidence of independent work:"), proposal["independentEvidence"], module["id"] + ":independentEvidence")
        labelled(choose("Beurteilungsvorschlag:", "Assessment proposal:"), proposal["assessmentProposal"], module["id"] + ":assessmentProposal")
        labelled(choose("Einbettung in den Arbeitsaufwand:", "Workload integration:"), proposal["workloadIntegration"], module["id"] + ":workloadIntegration")
        if proposal.get("topicRequired") or proposal.get("individualisationRequired"):
            flags = []
            if proposal.get("topicRequired"):
                flags.append(choose("Thema noch abzustimmen", "topic still to be agreed"))
            if proposal.get("individualisationRequired"):
                flags.append(choose("individuelle Konkretisierung erforderlich", "individual adaptation required"))
            paragraph(choose("Vor Durchführung: ", "Before implementation: ") + "; ".join(flags) + ".")
        labelled(choose("Offene Entscheidungen:", "Open decisions:"), proposal["openDecisions"], module["id"] + ":openDecisions")
        heading(choose("3  Lernziele und vorgeschlagene Nachweise", "3  Objectives and proposed evidence"), 2)
        for objective_index, objective in enumerate(proposal["objectives"], 1):
            goal = text(objective["text"])
            heading(f"{objective_index}. {goal}", 3)
            required.append((objective["id"] + ":text", goal))
            criterion = criteria[objective["criterionId"]]
            labelled(choose("Kriterium:", "Criterion:"), objective["criterionId"] + " · " + text(criterion["text"]))
            labelled(choose("Zielniveau / Kontext:", "Target level / context:"), ROMAN[objective["targetLevel"]] + " · " + "; ".join(contexts[id] for id in objective["contextIds"]))
            for key, de, en in (("task", "Aufgabe:", "Task:"), ("evidence", "Vorgeschlagener Nachweis:", "Proposed evidence:"), ("successCriteria", "Erfolgskriterium im Entwurf:", "Proposed success criterion:")):
                labelled(choose(de, en), objective[key], objective["id"] + ":" + key)
            marker = paragraph(f"[objective:{objective['id']}]")
            marker.runs[0].font.size = Pt(9)
            marker.runs[0].font.color.rgb = RGBColor.from_string(layout.GREY)
            marker.paragraph_format.space_after = Pt(7)
        spiral_start = len(document.paragraphs)
        heading(choose("4  Anschlüsse im Spiralcurriculum", "4  Links within the spiral curriculum"), 2)
        for direction, de, en in (("prior", "Anknüpfen an", "Build on"), ("next", "Später weiterführen", "Revisit later")):
            links = proposal["spiralLinks"][direction]
            heading(choose(de, en), 3)
            if not links:
                paragraph(choose("In diesem Entwurf sind keine konkreten Modulanschlüsse benannt.", "This proposal names no specific module links."))
            for link in links:
                neighbour = module_by_id[link["moduleId"]]
                p = paragraph("")
                internal_link(p, f"{neighbour['code']} · {text(neighbour['baseline']['title'])}", anchors[link["moduleId"]], layout.TEAL)
                p.paragraph_format.keep_with_next = True
                source_text(module["id"] + ":spiral:" + direction + ":" + link["moduleId"], link["reason"])

        # Keep this short, related navigation block together; avoid orphaned links.
        for p in document.paragraphs[spiral_start:-1]:
            p.paragraph_format.keep_with_next = True

    heading(choose("Quellenstand und Vollständigkeit", "Source record and completeness")).paragraph_format.page_break_before = True
    paragraph(choose("Dieser Band wurde aus denselben veröffentlichten Entwurfsdaten erzeugt, die auch dem Spiel zugrunde liegen. Die Originaldateien wurden nicht verändert.", "This volume was generated from the same public design data used by the game. Original files were not changed."))
    paragraph(f"Framework: {EXPECTED_VERSION} · {len(modules)} {choose('Einträge', 'entries')} · {EXPECTED_OBJECTIVES} {choose('Lernziele', 'objectives')}")
    for source in data["sources"]:
        if source.get("url"):
            p = paragraph("")
            layout.hyperlink(p, choose("Öffentliches Ausgangsspiel", "Public source game"), source["url"])
        if source.get("sourceDate"):
            paragraph(choose("Quellenstand: ", "Source date: ") + source["sourceDate"])
    paragraph("content/module-learning-design.json")
    paragraph("SHA256: " + sha256(SOURCE))
    paragraph("content/competency-framework.json")
    paragraph("SHA256: " + sha256(FRAMEWORK))
    paragraph(choose("Die 153 Lernziele sind vollständig enthalten. Wiederkehrende Kriterien bezeichnen dieselbe Kompetenz; sie begründen keine Mehrfachbewertung derselben Leistung. Fachliche Erprobung, verbindliche Modulbeschlüsse und eine psychometrische Validierung werden nicht behauptet.", "All 153 objectives are included. Recurring criteria refer to the same competency; they do not justify counting one performance several times. This document does not claim completed disciplinary piloting, binding module decisions or psychometric validation."))
    path = output_folder / f"Modulbeschreibungen-Psychologie-Zukunft-v2.1-{language.upper()}.docx"
    document.save(path)
    return verify_output(path, required, [m["id"] for m in modules], [o["id"] for m in modules for o in m["proposal"]["objectives"]])


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validate-only", action="store_true", help="Validate final source without writing documents.")
    parser.add_argument("--language", choices=("de", "en", "both"), default="both")
    args = parser.parse_args()
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    framework = json.loads(FRAMEWORK.read_text(encoding="utf-8"))
    validate_source(data, framework)
    if args.validate_only:
        print(json.dumps({"valid": True, "version": EXPECTED_VERSION, "modules": EXPECTED_ENTRIES, "objectives": EXPECTED_OBJECTIVES, "sourceSha256": sha256(SOURCE)}, ensure_ascii=False))
        return
    source_hash, framework_hash = sha256(SOURCE), sha256(FRAMEWORK)
    layout = load_layout()
    for language in ("de", "en") if args.language == "both" else (args.language,):
        if (sha256(SOURCE), sha256(FRAMEWORK)) != (source_hash, framework_hash):
            fail("Source changed during document generation; rerun from stable files.")
        result = build(language, data, framework, layout, ROOT.parent)
        print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.dont_write_bytecode = True
    main()
